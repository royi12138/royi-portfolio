from __future__ import annotations

import re
import shutil
import tempfile
from pathlib import Path
from typing import Any

import fitz
import pytesseract
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse
from starlette.background import BackgroundTask
from pydantic import BaseModel, Field
from PIL import Image, ImageOps

from app.ai_analysis import analyze_resume_text, generate_resume_revision, match_resume_to_jd, redact_sensitive_info

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
TMP_DIR = BASE_DIR.parent / "tmp"
TMP_DIR.mkdir(exist_ok=True)

MAX_BYTES = 10 * 1024 * 1024
MAX_ANALYSIS_CHARS = 20_000
SUPPORTED = {".pdf", ".docx", ".jpg", ".jpeg", ".png"}

app = FastAPI(title="AI 岗位匹配简历助手", version="0.7.3")

SECTION_ALIASES = {
    "basic_info": {"个人信息", "基本信息", "个人资料", "联系方式", "contact", "profile"},
    "education": {"教育经历", "教育背景", "学历", "教育", "education"},
    "work_experience": {"工作经历", "工作经验", "职业经历", "任职经历", "experience", "work experience", "employment"},
    "project_experience": {"项目经历", "项目经验", "项目", "projects", "project experience"},
    "skills": {"技能", "专业技能", "技能特长", "skills", "technical skills"},
    "self_intro": {"自我评价", "个人简介", "summary", "profile summary"},
}


class ResumeAnalysisRequest(BaseModel):
    resume_text: str = Field(min_length=30, max_length=MAX_ANALYSIS_CHARS)
    privacy_consent: bool = False


class ResumeRevisionRequest(BaseModel):
    resume_text: str = Field(min_length=30, max_length=MAX_ANALYSIS_CHARS)
    accepted_issues: list[dict[str, str]] = Field(default_factory=list, max_length=8)
    privacy_consent: bool = False
    facts: str = Field(default="", max_length=10_000)
    target_jd: str = Field(default="", max_length=10_000)


class ResumeMatchRequest(BaseModel):
    resume_text: str = Field(min_length=30, max_length=MAX_ANALYSIS_CHARS)
    jd_text: str = Field(min_length=30, max_length=10_000)
    privacy_consent: bool = False


class ResumeExportRequest(BaseModel):
    resume_text: str = Field(min_length=30, max_length=30_000)
    filename: str = Field(default="优化后简历", max_length=120)


def safe_filename(name: str) -> str:
    name = Path(name or "upload").name
    name = re.sub(r"[^\w.\-\u4e00-\u9fff ]+", "_", name)
    return name[:120] or "upload"


def normalize_text(text: str) -> str:
    text = text.replace("\u00a0", " ").replace("\r", "")
    lines = [re.sub(r"[ \t]+", " ", x).strip() for x in text.split("\n")]
    return "\n".join(x for x in lines if x)


def extract_pdf(path: Path) -> tuple[str, int]:
    doc = fitz.open(path)
    try:
        pages = doc.page_count
        parts = []
        for i, page in enumerate(doc):
            text = page.get_text("text") or ""
            if text.strip():
                parts.append(f"[第{i + 1}页]\n{text}")
        return normalize_text("\n\n".join(parts)), pages
    finally:
        doc.close()


def iter_docx_blocks(container: Any):
    """Yield paragraphs and tables in their real Word document order."""
    if hasattr(container, "element") and hasattr(container.element, "body"):
        element = container.element.body
    else:
        element = getattr(container, "_element", container)
    for child in element.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, container)
        elif child.tag.endswith("}tbl"):
            yield Table(child, container)


def extract_docx_cell(cell: Any) -> list[str]:
    parts: list[str] = []
    for block in iter_docx_blocks(cell):
        if isinstance(block, Paragraph):
            if block.text.strip():
                parts.append(block.text.strip())
        else:
            rows = []
            seen_cells: set[int] = set()
            for row in block.rows:
                cells = []
                for nested_cell in row.cells:
                    cell_id = id(nested_cell._tc)
                    if cell_id in seen_cells:
                        continue
                    seen_cells.add(cell_id)
                    nested_parts = extract_docx_cell(nested_cell)
                    if nested_parts:
                        cells.append("\n".join(nested_parts))
                if cells:
                    rows.append(" | ".join(cells))
            parts.extend(rows)
    return parts


def extract_docx(path: Path) -> tuple[str, int]:
    doc = Document(path)
    parts: list[str] = []
    seen_header_footer_parts: set[tuple[int, str]] = set()

    # Include text in page headers/footers (contact details are often placed there).
    for section in doc.sections:
        for attribute in (
            "first_page_header", "header", "even_page_header",
            "first_page_footer", "footer", "even_page_footer",
        ):
            container = getattr(section, attribute)
            key = (id(container.part), attribute.endswith("footer"))
            if key in seen_header_footer_parts:
                continue
            seen_header_footer_parts.add(key)
            for block in iter_docx_blocks(container):
                if isinstance(block, Paragraph):
                    if block.text.strip():
                        parts.append(block.text.strip())
                else:
                    for row in block.rows:
                        values = ["\n".join(extract_docx_cell(cell)) for cell in row.cells]
                        values = [value for value in values if value]
                        if values:
                            parts.append(" | ".join(values))
    body_paragraphs = []
    for block in iter_docx_blocks(doc):
        if isinstance(block, Paragraph):
            if block.text.strip():
                parts.append(block.text.strip())
                body_paragraphs.append(block)
        else:
            rows = []
            seen_cells: set[int] = set()
            for row in block.rows:
                cells = []
                for cell in row.cells:
                    cell_id = id(cell._tc)
                    if cell_id in seen_cells:
                        continue
                    seen_cells.add(cell_id)
                    cell_parts = extract_docx_cell(cell)
                    if cell_parts:
                        cells.append("\n".join(cell_parts))
                if cells:
                    rows.append(" | ".join(cells))
            parts.extend(rows)
    text = normalize_text("\n".join(parts))
    breaks = sum(paragraph._p.xml.count('w:type="page"') for paragraph in body_paragraphs)
    estimated = max(1, (len(text) + 1049) // 1050)
    pages = max(estimated, breaks + 1)
    return text, pages


def extract_image(path: Path) -> tuple[str, int]:
    image = Image.open(path)
    image = ImageOps.exif_transpose(image).convert("RGB")
    try:
        text = pytesseract.image_to_string(image, lang="chi_sim+eng", config="--psm 6")
    except pytesseract.TesseractError:
        text = pytesseract.image_to_string(image, config="--psm 6")
    return normalize_text(text), 1


def section_for_heading(line: str) -> str | None:
    normalized = re.sub(r"[\s:：|·•/\\]+", "", line).lower()
    for section, aliases in SECTION_ALIASES.items():
        for alias in aliases:
            alias_norm = re.sub(r"[\s:：|·•/\\]+", "", alias).lower()
            if normalized == alias_norm:
                return section
    if "工作" in line and ("经历" in line or "经验" in line):
        return "work_experience"
    if "项目" in line:
        return "project_experience"
    if "教育" in line or "学历" in line:
        return "education"
    if "技能" in line:
        return "skills"
    return None


def parse_sections(text: str) -> dict[str, list[str]]:
    data = {key: [] for key in SECTION_ALIASES}
    data["other"] = []
    current = "basic_info"
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        heading = section_for_heading(line)
        if heading:
            current = heading
        else:
            data[current].append(line)
    return data


def extract_basic_info(text: str, sections: dict[str, list[str]]) -> dict[str, Any]:
    email = re.findall(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", text)
    phone = re.findall(r"(?:\+?86[-\s]?)?1\d{10}", text)
    urls = re.findall(r"https?://\S+", text)
    return {
        "possible_email": email[:2],
        "possible_phone": phone[:2],
        "possible_urls": urls[:3],
        "raw_lines": sections.get("basic_info", [])[:12],
    }


def build_resume_json(filename: str, ext: str, text: str, pages: int) -> dict[str, Any]:
    sections = parse_sections(text)
    return {
        "document": {"filename": filename, "type": ext.lstrip("."), "pages": pages},
        "basic_info": extract_basic_info(text, sections),
        "education": sections["education"],
        "work_experience": sections["work_experience"],
        "project_experience": sections["project_experience"],
        "skills": sections["skills"],
        "self_intro": sections["self_intro"],
        "other": sections["other"],
        "raw_text": text,
        "parser": {"mode": "heuristic_mvp", "note": "已接入千问简历分析；语义抽取仍在后续完善。"},
    }


async def save_upload(file: UploadFile) -> Path:
    original = safe_filename(file.filename or "upload")
    suffix = Path(original).suffix.lower()
    if suffix not in SUPPORTED:
        raise HTTPException(400, "暂不支持该文件类型，请上传 PDF、DOCX、JPG 或 PNG。")

    tmp = Path(tempfile.mkdtemp(prefix="resume_", dir=TMP_DIR))
    target = tmp / original
    size = 0
    try:
        with target.open("wb") as out:
            while True:
                chunk = await file.read(1024 * 1024)
                if not chunk:
                    break
                size += len(chunk)
                if size > MAX_BYTES:
                    raise HTTPException(413, "文件不能超过 10MB。")
                out.write(chunk)
        return target
    except Exception:
        shutil.rmtree(tmp, ignore_errors=True)
        raise


def parse_file(path: Path) -> tuple[str, int]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix in {".jpg", ".jpeg", ".png"}:
        return extract_image(path)
    raise ValueError("unsupported file")


@app.get("/")
def index() -> FileResponse:
    response = FileResponse(STATIC_DIR / "index.html")
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate"
    response.headers["Pragma"] = "no-cache"
    return response


@app.get("/api/health")
def health() -> dict[str, str]:
    return {"status": "ok", "version": "0.7.3"}


@app.post("/api/resume/parse")
async def parse_resume(file: UploadFile = File(...)) -> dict[str, Any]:
    path = await save_upload(file)
    try:
        text, pages = parse_file(path)
        if pages > 3:
            raise HTTPException(400, "目前仅支持 3 页以内的简历，请压缩后重新上传。")
        if len(text.strip()) < 30:
            raise HTTPException(422, "没有识别到足够的简历文字，请上传更清晰的文件。")
        resume = build_resume_json(path.name, path.suffix, text, pages)
        return {"ok": True, "resume": resume}
    finally:
        shutil.rmtree(path.parent, ignore_errors=True)


@app.post("/api/jd/parse")
async def parse_jd(file: UploadFile = File(...)) -> dict[str, Any]:
    filename = safe_filename(file.filename or "job-description")
    suffix = Path(filename).suffix.lower()
    if suffix not in {".jpg", ".jpeg", ".png"}:
        raise HTTPException(400, "岗位信息请上传 JPG 或 PNG 图片。")
    content = await file.read(MAX_BYTES + 1)
    if len(content) > MAX_BYTES:
        raise HTTPException(413, "岗位文件不能超过 10MB。")
    directory = Path(tempfile.mkdtemp(prefix="jd_", dir=TMP_DIR))
    path = directory / filename
    try:
        path.write_bytes(content)
        text, _ = extract_image(path)
        if len(text.strip()) < 30:
            raise HTTPException(422, "没有识别到足够的岗位文字，请换一份更清晰的文件。")
        return {"ok": True, "filename": filename, "jd_text": text[:10_000]}
    finally:
        shutil.rmtree(directory, ignore_errors=True)


@app.post("/api/resume/analyze")
async def analyze_resume(request: ResumeAnalysisRequest) -> dict[str, Any]:
    if not request.privacy_consent:
        raise HTTPException(400, "请先确认隐私提示后再发送简历进行 AI 分析。")
    analysis = await analyze_resume_text(request.resume_text)
    return {"ok": True, "analysis": analysis}


@app.post("/api/resume/revise")
async def revise_resume(request: ResumeRevisionRequest) -> dict[str, Any]:
    if not request.privacy_consent:
        raise HTTPException(400, "请先确认隐私提示后再发送内容进行 AI 改写。")
    accepted = []
    for item in request.accepted_issues[:8]:
        quote = str(item.get("source_quote", ""))[:300]
        if quote and quote in redact_sensitive_info(request.resume_text):
            accepted.append({key: str(item.get(key, ""))[:600] for key in ("id", "section", "source_quote", "problem", "why_it_matters", "suggestion", "question")})
    if not accepted:
        raise HTTPException(400, "请接受至少一条带有简历原文引用的建议后再生成简历。")
    draft = await generate_resume_revision(request.resume_text, accepted, request.facts, request.target_jd)
    return {"ok": True, "revised_text": draft, "fact_protection": "只采用已接受建议和你填写的真实信息；请下载前逐项核对。", "stored": False}


@app.post("/api/resume/match")
async def match_resume(request: ResumeMatchRequest) -> dict[str, Any]:
    if not request.privacy_consent:
        raise HTTPException(400, "请先确认隐私提示后再发送简历和招聘信息进行匹配。")
    result = await match_resume_to_jd(request.resume_text, request.jd_text)
    return {"ok": True, "match": result}


@app.post("/api/resume/export")
async def export_resume(request: ResumeExportRequest) -> FileResponse:
    directory = Path(tempfile.mkdtemp(prefix="resume_export_", dir=TMP_DIR))
    try:
        filename = safe_filename(request.filename).rsplit(".", 1)[0] or "优化后简历"
        output = directory / (filename + ".docx")
        document = Document()
        document.add_heading(filename, level=1)
        for line in request.resume_text.splitlines():
            if line.strip():
                document.add_paragraph(line.strip())
        document.save(output)
        return FileResponse(
            output,
            filename=output.name,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            background=BackgroundTask(shutil.rmtree, directory, ignore_errors=True),
        )
    except Exception:
        shutil.rmtree(directory, ignore_errors=True)
        raise
