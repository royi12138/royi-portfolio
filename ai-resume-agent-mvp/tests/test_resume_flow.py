from app.ai_analysis import (apply_resume_replacements, build_match_messages, build_messages, build_revision_messages, build_completion_payload, numbers_are_supported, restore_sensitive_info, calculate_scores, extract_completion_text, extract_vision_text, safe_vision_error_message, parse_json_object, redact_sensitive_info, validate_analysis, validate_resume_sections)
from app.main import app, build_resume_json, extract_docx, parse_sections
from fastapi.testclient import TestClient
from docx import Document
from io import BytesIO
from PIL import Image
import pytest

client = TestClient(app)

SAMPLE_RESUME = "拥有多年互联网产品管理经验，负责需求规划、跨部门协作和产品交付。持续跟进用户反馈，推动项目按计划落地。"


def test_scores_are_explainable_and_bounded():
    score = calculate_scores("教育经历\n本科\n工作经历\n负责上线并提升用户数 20%", [])
    assert 0 <= score["overall"] <= 100
    assert len(score["dimensions"]) == 4
    assert all("reason" in item and item["weight"] > 0 for item in score["dimensions"])


def test_analysis_only_keeps_exact_source_quotes():
    source = "负责产品上线并提升用户数"
    result = validate_analysis(
        {"summary": "ok", "issues": [
            {"source_quote": source, "problem": "可补充结果"},
            {"source_quote": "没有出现在原文中的内容", "problem": "无依据"},
        ]},
        source,
    )
    assert len(result["issues"]) == 1
    assert result["issues"][0]["source_quote"] == source
    assert "scores" in result


def test_analysis_includes_safe_per_issue_rewrite_and_validates_numbers():
    source = "负责产品上线，使用Python整理需求"
    result = validate_analysis(
        {"issues": [
            {"source_quote": source, "problem": "表达不够具体", "suggestion": "突出原文中的行动",
             "suggested_text": "负责产品上线，使用Python整理需求"},
            {"source_quote": source, "problem": "表达不够具体", "suggested_text": "负责产品上线，新增用户1000人"},
        ]},
        source,
    )
    assert result["issues"][0]["suggested_text"] == source
    assert result["issues"][1]["suggested_text"] == ""


def test_analysis_prompt_requests_source_bounded_rewrite():
    messages = build_messages("负责产品上线")
    assert '"suggested_text"' in messages[0]["content"]
    assert "不增加角色" in messages[0]["content"]


def test_common_identifiers_are_redacted():
    redacted = redact_sensitive_info("邮箱 a.user@example.com，电话 13800138000")
    assert "a.user@example.com" not in redacted
    assert "13800138000" not in redacted


def test_resume_parser_keeps_sections_and_raw_text():
    raw = "个人信息\n张三\n教育经历\n本科\n工作经历\n负责产品上线\n技能\nPython"
    sections = parse_sections(raw)
    assert sections["education"] == ["本科"]
    assert sections["work_experience"] == ["负责产品上线"]
    resume = build_resume_json("resume.docx", ".docx", raw, 1)
    assert resume["raw_text"] == raw




def test_docx_extraction_preserves_order_and_reads_nested_tables_headers_and_footers(tmp_path):
    document = Document()
    document.sections[0].header.paragraphs[0].text = "张三｜电话 13800138000"
    document.sections[0].footer.paragraphs[0].text = "个人信息"
    document.add_paragraph("教育经历")
    outer = document.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = "本科｜某某大学｜2020年"
    nested = outer.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "项目经历"
    nested.cell(0, 0).add_paragraph("校园服务平台｜负责需求梳理")
    document.add_paragraph("持续负责后续维护")
    path = tmp_path / "synthetic_resume.docx"
    document.save(path)

    text, pages = extract_docx(path)

    assert text.index("教育经历") < text.index("本科") < text.index("项目经历")
    assert text.index("项目经历") < text.index("校园服务平台") < text.index("持续负责后续维护")
    assert "张三" in text and "13800138000" in text and "个人信息" in text
    sections = parse_sections(text)
    assert sections["education"] == ["本科｜某某大学｜2020年"]
    assert sections["project_experience"] == ["校园服务平台｜负责需求梳理", "持续负责后续维护"]
    assert pages >= 1


def test_analysis_requires_privacy_consent():
    response = client.post("/api/resume/analyze", json={"resume_text": SAMPLE_RESUME, "privacy_consent": False})
    assert response.status_code == 400


def test_revision_requires_privacy_consent():
    response = client.post("/api/resume/revise", json={
        "resume_text": SAMPLE_RESUME,
        "accepted_issues": [{"source_quote": "拥有多年互联网产品管理经验", "suggestion": "优化表达"}],
        "privacy_consent": False,
    })
    assert response.status_code == 400


def test_jd_match_requires_privacy_consent():
    response = client.post("/api/resume/match", json={
        "resume_text": SAMPLE_RESUME,
        "jd_text": "招聘高级产品经理，负责产品规划、用户调研、需求分析、跨团队沟通与项目交付。",
        "privacy_consent": False,
    })
    assert response.status_code == 400


def test_docx_export_returns_word_document():
    response = client.post("/api/resume/export", json={"resume_text": SAMPLE_RESUME, "filename": "测试简历"})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert response.content.startswith(b"PK")


def test_homepage_prevents_stale_consent_script():
    response = client.get("/")
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store, no-cache, must-revalidate"
    assert "privacy_consent:document.getElementById('consent').checked" in response.text


def test_jd_image_upload_extracts_text(monkeypatch):
    jd = "招聘产品经理，负责产品规划、需求分析、用户研究、跨团队协作、项目推进与产品上线交付。"
    async def fake_vision(content, mime_type):
        assert content.startswith(b"\x89PNG")
        assert mime_type == "image/png"
        return jd
    monkeypatch.setattr("app.main.extract_job_text_from_image", fake_vision)
    image = BytesIO()
    Image.new("RGB", (40, 30), "white").save(image, format="PNG")
    response = client.post(
        "/api/jd/parse",
        files={"file": ("岗位.png", image.getvalue(), "image/png")},
    )
    assert response.status_code == 200
    assert response.json()["jd_text"] == jd


def test_jd_text_upload_is_rejected():
    response = client.post(
        "/api/jd/parse",
        files={"file": ("岗位.txt", b"job description long enough to pass", "text/plain")},
    )
    assert response.status_code == 400

def test_analysis_page_reveals_results_and_contains_comparison_views():
    response = client.get("/")
    assert response.status_code == 200
    assert "analysisResult').classList.remove('hidden')" in response.text
    assert 'id="beforeResume"' in response.text
    assert 'id="jdBeforeText"' in response.text
    assert "async function generateJDRevision()" in response.text
    assert "function renderParsedResume(resume)" in response.text
    assert "简历识别与 AI 诊断" not in response.text
    assert "识别到的简历信息（初步）" not in response.text
    assert 'id="scoreResults" class="score-card"' in response.text
    assert "score-radar" in response.text
    assert "需要你补充的真实信息" not in response.text
    assert "只填写可确认的真实信息" not in response.text
    assert "function readApiJson(response,failureMessage)" in response.text
    assert "function renderComparison(beforeId,afterId,outputId)" in response.text
    assert "diff-change" in response.text
    assert 'accept=".jpg,.jpeg,.png,image/jpeg,image/png" multiple' in response.text
    assert "function renderIssueComparison(before,after)" in response.text
    assert "async function parseSelectedJDImages()" in response.text
    assert "parseJDFile" not in response.text
    assert "开始 AI 简历诊断" in response.text
    assert response.text.index('id="consent"') < response.text.index('id="parseBtn"')
    assert "await analyzeResume(button)" in response.text
    assert "issue-added" in response.text
    assert "--ink:#182a27" in response.text
    assert ".hero{margin:0 0 20px" in response.text
    assert "@media(max-width:640px)" in response.text
    assert "MVP 0.5" not in response.text


def test_health_version_is_current():
    response = client.get("/api/health")
    assert response.json()["version"] == "0.7.6"



def test_completion_handles_text_blocks_and_glm_payload():
    response = {"choices": [{"message": {"content": [{"type": "text", "text": "ok"}, {"content": " done"}]}}]}
    assert extract_completion_text(response) == "ok\n done"
    payload = build_completion_payload("glm-5.3-flash", [], json_mode=True)
    assert payload["thinking"] == {"type": "enabled"}
    assert payload["reasoning_effort"] == "low"
    assert payload["clear_thinking"] is True
    assert payload["response_format"] == {"type": "json_object"}




def test_provider_error_message_is_sanitized():
    message = safe_vision_error_message(
        {"code": 1211, "message": "url error; email user@example.com; Authorization: Bearer abc123"}
    )
    assert "url error" in message
    assert "user@example.com" not in message
    assert "abc123" not in message


def test_extract_vision_text_reads_dashscope_native_response():
    response = {
        "output": {
            "choices": [{
                "message": {
                    "content": [{"text": "岗位职责一"}, {"text": "岗位要求二"}]
                }
            }]
        }
    }
    assert extract_vision_text(response) == "岗位职责一 岗位要求二"


def test_match_json_parser_accepts_fenced_or_prefixed_json():
    assert parse_json_object("```json\n{\"summary\":\"ok\"}\n```") == {"summary": "ok"}
    assert parse_json_object("回答：{\"summary\":\"ok\"}完成") == {"summary": "ok"}


def test_structured_resume_summary_requires_exact_source_and_supported_numbers():
    source = "工作业绩：注册用户增长2000+，转化率提升12%。"
    result = validate_resume_sections(
        {
            "work_experience": [
                {"summary": "• 注册用户增长 2000+；转化率提升12%。", "source_quote": source},
                {"summary": "新增营收999万元。", "source_quote": source},
                {"summary": "完全无关的总结", "source_quote": "不存在于原文"},
            ]
        },
        source,
    )
    assert result["work_experience"] == ["注册用户增长 2000+；转化率提升12%。"]
    assert result["education"] == []


def test_analysis_returns_clean_structured_resume_sections():
    source = "教育经历：本科，信息管理专业。"
    result = validate_analysis({
        "summary": "清晰",
        "resume_sections": {
            "education": [{"summary": "• 本科，信息管理专业。", "source_quote": source}]
        },
        "issues": [],
    }, source)
    assert result["resume_sections"]["education"] == ["本科，信息管理专业。"]

def test_advisor_prompts_require_specific_evidence_and_usable_rewrites():
    analysis_messages = build_messages("负责产品规划与上线，推动用户反馈闭环。")
    analysis_prompt = analysis_messages[0]["content"]
    assert "可直接替换文案" in analysis_prompt
    assert "不能添加原文没有的个人角色、决策、方法、原因" in analysis_prompt
    assert "按影响从高到低排序" in analysis_prompt
    assert "不要为了“有问题可提”而挑小毛病" in analysis_prompt

    revision_messages = build_revision_messages(
        "工作经历\n负责产品规划与上线",
        [{"section": "工作经历", "source_quote": "负责产品规划与上线", "suggestion": "突出行动和结果"}],
        "用户确认上线范围为3个城市",
        "招聘产品经理，负责产品规划",
    )
    revision_system, revision_user = revision_messages[0]["content"], revision_messages[1]["content"]
    assert "其余简历由程序原样保留" in revision_system
    assert "用户确认的真实补充事实" in revision_user
    assert "招聘产品经理" in revision_user
    assert "3个城市" in revision_user
    assert "不要编造基本信息、联系方式、技能栏目" in revision_system

    match_messages = build_match_messages("负责产品规划", "招聘产品经理，要求数据分析")
    match_prompt = match_messages[0]["content"]
    assert "核心职责、硬性条件、优先条件" in match_prompt
    assert "未提及不等于候选人没有该能力" in match_prompt
    assert "可直接替换文案只能重组 source_quote 明确的事实" in match_prompt

def test_analysis_filters_unsupported_numbers_from_model_claims():
    source = "会员续费率从68%提升至76%，月活会员增加1.2万人。"
    result = validate_analysis(
        {
            "summary": "续费率提升至80%。",
            "strengths": ["续费率提升至76%。", "用户规模达到2万人。"],
            "issues": [
                {
                    "source_quote": source,
                    "problem": "成果数据清楚。",
                    "why_it_matters": "便于评估业务结果。",
                    "suggestion": "保留续费率从68%提升至76%的表述。",
                },
                {
                    "source_quote": source,
                    "problem": "数据不足。",
                    "why_it_matters": "缺少业务数据。",
                    "suggestion": "改成续费率提升至80%。",
                },
            ],
        },
        source,
    )
    assert result["summary"] == "分析已完成，请逐条核对建议是否符合你的真实经历。"
    assert result["strengths"] == ["续费率提升至76%。"]
    assert len(result["issues"]) == 1
    assert numbers_are_supported("月活会员增加1.2万人", source)
    assert not numbers_are_supported("月活会员增加2万人", source)

def test_revision_only_replaces_accepted_quote_and_restores_redacted_contact():
    resume = "邮箱：candidate@example.com\n工作经历\n负责产品规划与上线。\n项目经历\n保留此段不变。电话13800138000"
    accepted = [{
        "id": "I1",
        "source_quote": "负责产品规划与上线。",
        "suggestion": "调整语序",
    }]
    result = {
        "replacements": [{
            "id": "I1",
            "source_quote": "负责产品规划与上线。",
            "replacement": "负责产品上线与规划。",
        }]
    }
    revised = apply_resume_replacements(resume, accepted, result)
    assert "负责产品上线与规划。" in revised
    assert "保留此段不变。" in revised
    assert "candidate@example.com" in revised
    assert "13800138000" in revised
    assert "邮箱已隐藏" not in revised


def test_revision_rejects_unanchored_or_unsupported_numeric_rewrites():
    resume = "工作经历\n续费率提升至76%。\n项目经历\n保留。"
    accepted = [{"id": "I1", "source_quote": "续费率提升至76%。"}]
    with pytest.raises(Exception) as err:
        apply_resume_replacements(
            resume,
            accepted,
            {"replacements": [{
                "id": "I1",
                "source_quote": "续费率提升至76%。",
                "replacement": "续费率提升至80%。",
            }]},
        )
    assert getattr(err.value, "status_code", None) == 422

    with pytest.raises(Exception) as err:
        apply_resume_replacements(
            resume,
            accepted,
            {"replacements": [{
                "id": "I1",
                "source_quote": "续费率提升至80%。",
                "replacement": "续费率提升至80%。",
            }]},
        )
    assert getattr(err.value, "status_code", None) == 502

